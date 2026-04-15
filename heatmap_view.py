import streamlit as st
import pandas as pd
import json
import io
import matplotlib
import ai_insights as ai_mod
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.colors as mcolors
from matplotlib.backends.backend_agg import FigureCanvasAgg
import numpy as np
from investigation_standards import InvestigationStandards
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def render_heatmap_view(db, selected_patient, user_id, edit_df, key_prefix="manage", editable=True):
    patient_profile = db.get_patient_profile(selected_patient, user_id)
    patient_gender = patient_profile.get('gender', '').lower() if patient_profile else ''

    investigation_standards = InvestigationStandards(db)
    all_standards = investigation_standards.get_all_standards()

    test_info = {}
    categories_dict = {}

    for standard in all_standards:
        if isinstance(standard, dict):
            category = standard.get('category', 'Other')
            test_name = standard.get('standard_name', '').strip()
            aliases = standard.get('aliases', '')

            if patient_gender == 'male':
                range_min = standard.get('male_range_min')
                range_max = standard.get('male_range_max')
                unit = standard.get('male_unit', '')
            elif patient_gender == 'female':
                range_min = standard.get('female_range_min')
                range_max = standard.get('female_range_max')
                unit = standard.get('female_unit', '')
            else:
                range_min = standard.get('male_range_min')
                range_max = standard.get('male_range_max')
                unit = standard.get('male_unit', '')
        else:
            if len(standard) > 4:
                category = standard[4] if standard[4] else 'Other'
                test_name = standard[1] if len(standard) > 1 else ''
                aliases = standard[5] if len(standard) > 5 and standard[5] else ""
                range_min = standard[6] if len(standard) > 6 else None
                range_max = standard[7] if len(standard) > 7 else None
                unit = standard[8] if len(standard) > 8 else ''
            else:
                continue

        if category and test_name:
            if category not in categories_dict:
                categories_dict[category] = []
            categories_dict[category].append(test_name)

            test_info[test_name] = {
                'category': category,
                'min': range_min,
                'max': range_max,
                'unit': unit
            }

            if aliases:
                for alias in aliases.split(','):
                    alias = alias.strip()
                    if alias:
                        test_info[alias] = test_info[test_name]
                        if alias not in categories_dict[category]:
                            categories_dict[category].append(alias)

    col1, col2 = st.columns([1, 1])
    with col1:
        date_format = st.selectbox(
            "Date Format:",
            options=["UK (DD/MM/YYYY)", "US (MM/DD/YYYY)"],
            key=f"{key_prefix}_heatmap_date_format"
        )
    with col2:
        sort_tests = st.selectbox(
            "Sort Tests:",
            options=["Alphabetically", "By Category"],
            key=f"{key_prefix}_heatmap_sort_tests"
        )

    st.divider()

    all_patient_tests = edit_df['Test Type'].unique().tolist()

    patient_tests_by_category = {}
    uncategorized_tests = []
    for test in all_patient_tests:
        info = test_info.get(test)
        if info:
            cat = info['category']
            if cat not in patient_tests_by_category:
                patient_tests_by_category[cat] = []
            patient_tests_by_category[cat].append(test)
        else:
            uncategorized_tests.append(test)

    if uncategorized_tests:
        patient_tests_by_category['Other'] = uncategorized_tests

    filter_key = f"{key_prefix}_heatmap_filter_{selected_patient}"
    if filter_key not in st.session_state:
        saved_tests_json = db.get_user_preference(user_id, selected_patient, f'{key_prefix}_heatmap_selected_tests')
        if not saved_tests_json and key_prefix == 'manage':
            saved_tests_json = db.get_user_preference(user_id, selected_patient, 'heatmap_selected_tests')
        if saved_tests_json:
            try:
                saved_tests = json.loads(saved_tests_json)
                st.session_state[filter_key] = {'groups': {}, 'tests': {t: True for t in saved_tests if t in all_patient_tests}}
            except:
                st.session_state[filter_key] = {'groups': {}, 'tests': {t: True for t in all_patient_tests}}
        else:
            st.session_state[filter_key] = {'groups': {}, 'tests': {t: True for t in all_patient_tests}}

    st.write("**Filter by Test Groups:**")

    selected_tests = []

    with st.expander("📋 Select Test Groups and Tests", expanded=True):
        btn_col1, btn_col2, btn_col3 = st.columns([1, 1, 2])
        with btn_col1:
            if st.button("✅ Select All", key=f"{key_prefix}_select_all_tests"):
                for cat in patient_tests_by_category:
                    st.session_state[filter_key]['groups'][cat] = True
                    for test in patient_tests_by_category[cat]:
                        st.session_state[filter_key]['tests'][test] = True
                st.rerun()
        with btn_col2:
            if st.button("❌ Deselect All", key=f"{key_prefix}_deselect_all_tests"):
                st.session_state[filter_key]['groups'] = {}
                st.session_state[filter_key]['tests'] = {}
                st.rerun()
        with btn_col3:
            if st.button("💾 Save Filter Selection", key=f"{key_prefix}_save_filter_selection"):
                tests_to_save = [t for t, selected in st.session_state[filter_key]['tests'].items() if selected]
                if db.save_user_preference(user_id, selected_patient, f'{key_prefix}_heatmap_selected_tests', json.dumps(tests_to_save)):
                    st.success("Filter selection saved! It will be restored next time.")
                else:
                    st.error("Failed to save filter selection.")

        for category in sorted(patient_tests_by_category.keys()):
            tests_in_cat = sorted(patient_tests_by_category[category], key=lambda x: x.lower())

            all_selected = all(st.session_state[filter_key]['tests'].get(t, False) for t in tests_in_cat)
            any_selected = any(st.session_state[filter_key]['tests'].get(t, False) for t in tests_in_cat)

            group_checked = st.checkbox(
                f"**{category}** ({len(tests_in_cat)} tests)",
                value=all_selected,
                key=f"{key_prefix}_group_{category}"
            )

            prev_group_state = st.session_state[filter_key]['groups'].get(category, all_selected)
            if group_checked != prev_group_state:
                st.session_state[filter_key]['groups'][category] = group_checked
                for test in tests_in_cat:
                    st.session_state[filter_key]['tests'][test] = group_checked
                st.rerun()

            with st.container():
                cols = st.columns([0.05, 0.95])
                with cols[1]:
                    for test in tests_in_cat:
                        test_checked = st.checkbox(
                            test,
                            value=st.session_state[filter_key]['tests'].get(test, False),
                            key=f"{key_prefix}_test_{category}_{test}"
                        )
                        if test_checked:
                            selected_tests.append(test)
                        st.session_state[filter_key]['tests'][test] = test_checked

    if not selected_tests:
        st.warning("No tests selected. Please select at least one test to view the heat map.")
        selected_tests = []

    heatmap_df = edit_df[edit_df['Test Type'].isin(selected_tests)].copy()
    heatmap_df['Date_dt'] = pd.to_datetime(heatmap_df['Date'])

    if date_format == "UK (DD/MM/YYYY)":
        heatmap_df['Date_Formatted'] = heatmap_df['Date_dt'].dt.strftime('%d/%m/%Y')
    else:
        heatmap_df['Date_Formatted'] = heatmap_df['Date_dt'].dt.strftime('%m/%d/%Y')

    if heatmap_df.empty:
        st.warning("No data available for the selected tests.")
    else:
        unique_dates = heatmap_df.sort_values('Date_dt', ascending=False)['Date_Formatted'].unique().tolist()

        unique_tests = heatmap_df['Test Type'].unique().tolist()

        if sort_tests == "Alphabetically":
            unique_tests = sorted(unique_tests, key=lambda x: x.lower())
        else:
            def get_category_sort_key(test):
                info = test_info.get(test)
                cat = info['category'] if info else 'ZZZ_Other'
                return (cat, test.lower())
            unique_tests = sorted(unique_tests, key=get_category_sort_key)

        pivot_data = {}
        reference_ranges = {}
        value_ids = {}
        units_map = {}

        for test in unique_tests:
            pivot_data[test] = {}
            value_ids[test] = {}
            test_rows = heatmap_df[heatmap_df['Test Type'] == test]

            info = test_info.get(test)
            if info and info.get('min') is not None and info.get('max') is not None:
                try:
                    reference_ranges[test] = {'min': float(info['min']), 'max': float(info['max'])}
                    units_map[test] = info.get('unit', '')
                except (ValueError, TypeError):
                    reference_ranges[test] = None
                    units_map[test] = ''
            else:
                std_name, std_match = investigation_standards.standardize_test_name(test, source_text=test)
                if std_match:
                    if patient_gender == 'male':
                        range_min = std_match.get('male_range_min')
                        range_max = std_match.get('male_range_max')
                        unit = std_match.get('male_unit', '')
                    elif patient_gender == 'female':
                        range_min = std_match.get('female_range_min')
                        range_max = std_match.get('female_range_max')
                        unit = std_match.get('female_unit', '')
                    else:
                        range_min = std_match.get('male_range_min')
                        range_max = std_match.get('male_range_max')
                        unit = std_match.get('male_unit', '')

                    if range_min is not None and range_max is not None:
                        try:
                            reference_ranges[test] = {'min': float(range_min), 'max': float(range_max)}
                            units_map[test] = unit
                        except (ValueError, TypeError):
                            reference_ranges[test] = None
                            units_map[test] = ''
                    else:
                        reference_ranges[test] = None
                        units_map[test] = ''
                else:
                    if not test_rows.empty:
                        ref_str = test_rows.iloc[0]['Reference Range']
                        units_map[test] = test_rows.iloc[0]['Unit'] or ''
                        if ref_str and '-' in str(ref_str):
                            try:
                                parts = str(ref_str).replace(' ', '').split('-')
                                if len(parts) >= 2:
                                    low = float(parts[0])
                                    high_part = parts[1].split()[0] if ' ' in parts[1] else parts[1]
                                    high = float(''.join(c for c in high_part if c.isdigit() or c == '.'))
                                    reference_ranges[test] = {'min': low, 'max': high}
                                else:
                                    reference_ranges[test] = None
                            except:
                                reference_ranges[test] = None
                        else:
                            reference_ranges[test] = None
                    else:
                        reference_ranges[test] = None
                        units_map[test] = ''

            for date in unique_dates:
                date_rows = test_rows[test_rows['Date_Formatted'] == date]
                if not date_rows.empty:
                    pivot_data[test][date] = date_rows.iloc[0]['Value']
                    value_ids[test][date] = date_rows.iloc[0]['ID']
                else:
                    pivot_data[test][date] = None
                    value_ids[test][date] = None

        pivot_df = pd.DataFrame(pivot_data).T
        pivot_df.index.name = 'Test Name'

        def _text_color_for_bg(r, g, b):
            brightness = 0.299 * r + 0.587 * g + 0.114 * b
            return "white" if brightness < 140 else "black"

        def get_deviation_color(value, ref_info):
            if value is None or pd.isna(value):
                return ""

            try:
                value = float(value)
            except (ValueError, TypeError):
                return ""

            low, high = None, None

            if isinstance(ref_info, dict) and ref_info:
                low = ref_info.get('min')
                high = ref_info.get('max')

            if low is None or high is None:
                return ""

            try:
                low, high = float(low), float(high)
            except (ValueError, TypeError):
                return ""

            if high <= low:
                return ""

            range_size = high - low

            if low <= value <= high:
                mid = (low + high) / 2
                deviation = abs(value - mid) / (range_size / 2)
                deviation = min(max(deviation, 0), 1)
                r = int(144 - deviation * 110)
                g = int(238 - deviation * 99)
                b = int(144 - deviation * 110)
                fc = _text_color_for_bg(r, g, b)
                return f"background-color: rgb({r}, {g}, {b}); color: {fc}"
            elif value < low:
                deviation = (low - value) / range_size
                deviation = min(max(deviation, 0), 2)
                intensity = min(deviation / 2, 1)
                r = int(173 - intensity * 173)
                g = int(216 - intensity * 216)
                b = int(230 - intensity * 91)
                fc = _text_color_for_bg(r, g, b)
                return f"background-color: rgb({r}, {g}, {b}); color: {fc}"
            else:
                deviation = (value - high) / range_size
                deviation = min(max(deviation, 0), 2)
                intensity = min(deviation / 2, 1)
                r = int(255 - intensity * 116)
                g = int(204 - intensity * 204)
                b = int(203 - intensity * 203)
                fc = _text_color_for_bg(r, g, b)
                return f"background-color: rgb({r}, {g}, {b}); color: {fc}"

        st.markdown("""
        <style>
        .legend-container {
            display: flex;
            align-items: center;
            gap: 20px;
            padding: 10px;
            background: #f8f9fa;
            border-radius: 8px;
            margin-bottom: 10px;
        }
        .legend-item {
            display: flex;
            align-items: center;
            gap: 8px;
        }
        .legend-gradient {
            width: 80px;
            height: 20px;
            border-radius: 4px;
            border: 1px solid #ccc;
        }
        .blue-gradient {
            background: linear-gradient(to right, #00008B, #ADD8E6);
        }
        .green-gradient {
            background: linear-gradient(to right, #228B22, #90EE90, #228B22);
        }
        .red-gradient {
            background: linear-gradient(to right, #FFCCCB, #8B0000);
        }
        </style>
        <div class="legend-container">
            <strong>Color Legend:</strong>
            <div class="legend-item">
                <div class="legend-gradient blue-gradient"></div>
                <span>Below Normal</span>
            </div>
            <div class="legend-item">
                <div class="legend-gradient green-gradient"></div>
                <span>Normal Range</span>
            </div>
            <div class="legend-item">
                <div class="legend-gradient red-gradient"></div>
                <span>Above Normal</span>
            </div>
        </div>
        """, unsafe_allow_html=True)

        patient_id_for_events = db.get_patient_id(selected_patient, user_id)
        patient_events = []
        event_dates_map = {}
        if patient_id_for_events:
            patient_events = db.get_patient_events(patient_id_for_events, user_id)
            for ev in patient_events:
                ev_date = ev.get('event_date', '')
                if ev_date:
                    try:
                        ev_dt = pd.to_datetime(ev_date).normalize()
                        if date_format == "UK (DD/MM/YYYY)":
                            ev_formatted = ev_dt.strftime('%d/%m/%Y')
                        else:
                            ev_formatted = ev_dt.strftime('%m/%d/%Y')
                        if ev_formatted not in event_dates_map:
                            event_dates_map[ev_formatted] = []
                        event_dates_map[ev_formatted].append(ev.get('event_name', ''))
                    except:
                        pass

        show_events = False
        if patient_events:
            show_events = st.toggle("Show Patient Events in Heat Map", value=False, key=f"{key_prefix}_show_events_toggle")

        if patient_gender:
            st.info(f"Using {patient_gender.title()} reference ranges")
        st.write(f"**{len(unique_tests)} tests × {len(unique_dates)} dates** (Latest date first)")

        stats_data = {}
        for test in unique_tests:
            values = [v for v in pivot_data[test].values() if v is not None and not pd.isna(v)]
            if values:
                try:
                    float_values = [float(v) for v in values]
                    stats_data[test] = {
                        'Count': len(float_values),
                        'Sum': round(sum(float_values), 2),
                        'Average': round(sum(float_values) / len(float_values), 2)
                    }
                except:
                    stats_data[test] = {'Count': 0, 'Sum': 0, 'Average': 0}
            else:
                stats_data[test] = {'Count': 0, 'Sum': 0, 'Average': 0}

        st.markdown("**Statistics Summary:**")
        stats_cols = st.columns(3)
        total_count = sum(s['Count'] for s in stats_data.values())
        total_sum = sum(s['Sum'] for s in stats_data.values())
        overall_avg = total_sum / total_count if total_count > 0 else 0
        with stats_cols[0]:
            st.metric("Total Values", total_count)
        with stats_cols[1]:
            st.metric("Total Sum", f"{total_sum:.2f}")
        with stats_cols[2]:
            st.metric("Overall Average", f"{overall_avg:.2f}")

        range_column = {}
        for test in unique_tests:
            ref = reference_ranges.get(test)
            if ref and isinstance(ref, dict):
                unit = units_map.get(test, '')
                range_column[test] = f"{ref['min']}-{ref['max']} {unit}".strip()
            else:
                range_column[test] = "N/A"

        enhanced_pivot_df = pivot_df.copy()
        enhanced_pivot_df.insert(0, 'Range', [range_column.get(t, 'N/A') for t in enhanced_pivot_df.index])
        enhanced_pivot_df['Avg'] = [stats_data.get(t, {}).get('Average', 0) for t in enhanced_pivot_df.index]
        enhanced_pivot_df['Count'] = [stats_data.get(t, {}).get('Count', 0) for t in enhanced_pivot_df.index]
        enhanced_pivot_df['Sum'] = [stats_data.get(t, {}).get('Sum', 0) for t in enhanced_pivot_df.index]

        summary_row = {'Range': 'SUMMARY'}
        for date_col in unique_dates:
            date_values = []
            for test in unique_tests:
                val = pivot_data.get(test, {}).get(date_col)
                if val is not None and not pd.isna(val):
                    try:
                        date_values.append(float(val))
                    except:
                        pass
            if date_values:
                summary_row[date_col] = round(sum(date_values) / len(date_values), 2)
            else:
                summary_row[date_col] = None
        summary_row['Avg'] = overall_avg
        summary_row['Count'] = total_count
        summary_row['Sum'] = total_sum

        summary_df = pd.DataFrame([summary_row], index=['Average Per Date'])
        enhanced_pivot_df = pd.concat([enhanced_pivot_df, summary_df])

        event_columns = set()
        if show_events and event_dates_map:
            all_date_objs = []
            for d in unique_dates:
                try:
                    if date_format == "UK (DD/MM/YYYY)":
                        dt = pd.to_datetime(d, format='%d/%m/%Y')
                    else:
                        dt = pd.to_datetime(d, format='%m/%d/%Y')
                    all_date_objs.append((dt, d, 'data'))
                except:
                    all_date_objs.append((pd.Timestamp.min, d, 'data'))

            for ev_formatted, ev_names in event_dates_map.items():
                if ev_formatted not in unique_dates:
                    try:
                        if date_format == "UK (DD/MM/YYYY)":
                            ev_dt = pd.to_datetime(ev_formatted, format='%d/%m/%Y')
                        else:
                            ev_dt = pd.to_datetime(ev_formatted, format='%m/%d/%Y')
                        all_date_objs.append((ev_dt, ev_formatted, 'event'))
                    except:
                        all_date_objs.append((pd.Timestamp.min, ev_formatted, 'event'))

            all_date_objs.sort(key=lambda x: x[0], reverse=True)

            event_col_labels = {}
            for ev_formatted, ev_names in event_dates_map.items():
                event_label = f"📌 {ev_formatted}"
                event_col_labels[ev_formatted] = event_label

            merged_columns = []
            for dt_obj, formatted, dtype in all_date_objs:
                if dtype == 'event' and formatted not in unique_dates:
                    label = event_col_labels.get(formatted, f"📌 {formatted}")
                    merged_columns.append(label)
                    event_columns.add(label)
                    event_name_str = ' | '.join(event_dates_map.get(formatted, []))
                    enhanced_pivot_df[label] = event_name_str
                else:
                    merged_columns.append(formatted)
                    if formatted in event_dates_map:
                        label = event_col_labels.get(formatted, f"📌 {formatted}")
                        if label not in merged_columns:
                            idx = merged_columns.index(formatted)
                            merged_columns.insert(idx + 1, label)
                            event_columns.add(label)
                            event_name_str = ' | '.join(event_dates_map.get(formatted, []))
                            enhanced_pivot_df[label] = event_name_str

            cols_order = ['Range'] + merged_columns + ['Avg', 'Count', 'Sum']
            existing_cols = [c for c in cols_order if c in enhanced_pivot_df.columns]
            enhanced_pivot_df = enhanced_pivot_df[existing_cols]

        date_columns = [c for c in enhanced_pivot_df.columns if c in unique_dates]

        def apply_row_style(row):
            test_name = row.name
            if test_name == 'Average Per Date':
                return ['background-color: #e8e8e8; font-weight: bold'] * len(row)
            ref_info = reference_ranges.get(test_name, "")
            styles = []
            for col in row.index:
                if col in event_columns:
                    styles.append('background-color: #FFF3CD; font-weight: bold; writing-mode: vertical-rl; text-orientation: mixed; font-size: 10px; color: #856404;')
                elif col in date_columns:
                    color = get_deviation_color(row[col], ref_info)
                    styles.append(color if color else "")
                else:
                    styles.append("")
            return styles

        styled_display = enhanced_pivot_df.style.apply(apply_row_style, axis=1)
        styled_display = styled_display.format(precision=2, na_rep="")

        st.markdown("""
        <style>
        /* Auto-fit column widths for st.dataframe */
        div[data-testid="stDataFrame"] > div {
            overflow-x: auto !important;
        }
        div[data-testid="stDataFrame"] table {
            width: auto !important;
            table-layout: auto !important;
        }
        div[data-testid="stDataFrame"] table th,
        div[data-testid="stDataFrame"] table td {
            white-space: nowrap !important;
            padding: 4px 8px !important;
        }
        /* Sticky header for st.dataframe */
        div[data-testid="stDataFrame"] table thead tr th {
            position: sticky !important;
            top: 0 !important;
            background-color: #f0f2f6 !important;
            z-index: 100 !important;
        }
        /* Freeze first column for st.dataframe */
        div[data-testid="stDataFrame"] table tbody tr td:first-child,
        div[data-testid="stDataFrame"] table thead tr th:first-child {
            position: sticky !important;
            left: 0 !important;
            background-color: #f0f2f6 !important;
            z-index: 99 !important;
        }
        div[data-testid="stDataFrame"] table thead tr th:first-child {
            z-index: 101 !important;
        }
        
        /* Auto-fit column widths for st.data_editor */
        div[data-testid="stDataEditor"] > div {
            overflow-x: auto !important;
        }
        div[data-testid="stDataEditor"] table {
            width: auto !important;
            table-layout: auto !important;
        }
        div[data-testid="stDataEditor"] table th,
        div[data-testid="stDataEditor"] table td {
            white-space: nowrap !important;
            padding: 4px 8px !important;
        }
        /* Sticky header for st.data_editor */
        div[data-testid="stDataEditor"] table thead tr th {
            position: sticky !important;
            top: 0 !important;
            background-color: #f0f2f6 !important;
            z-index: 100 !important;
        }
        /* Freeze first column for st.data_editor */
        div[data-testid="stDataEditor"] table tbody tr td:first-child,
        div[data-testid="stDataEditor"] table thead tr th:first-child {
            position: sticky !important;
            left: 0 !important;
            background-color: #f0f2f6 !important;
            z-index: 99 !important;
        }
        div[data-testid="stDataEditor"] table thead tr th:first-child {
            z-index: 101 !important;
        }
        </style>
        """, unsafe_allow_html=True)
        st.caption("Tip: Scroll horizontally to see all dates. Test name column and header row stay fixed when scrolling.")

        with st.expander("📊 Heat Map View (Color-Coded)", expanded=True):
            st.dataframe(styled_display, use_container_width=False, height=min(600, len(unique_tests) * 35 + 80))

        # ── AI / Web Search Analysis Panel ───────────────────────────────────
        with st.expander("🔍 Analyse a Test Result — Web Search or AI", expanded=False):
            stored_keys = db.get_all_ai_api_keys(user_id)
            configured_svcs = [s for s in ai_mod.SERVICES if stored_keys.get(s)]

            col_t, col_d = st.columns(2)
            with col_t:
                sel_test_ai = st.selectbox(
                    "Test", options=unique_tests, key=f"{key_prefix}_ai_test"
                )
            with col_d:
                sel_date_ai = st.selectbox(
                    "Date", options=unique_dates, key=f"{key_prefix}_ai_date"
                )

            # Look up the value for the selected cell
            cell_value = None
            if sel_test_ai and sel_date_ai:
                raw = pivot_data.get(sel_test_ai, {}).get(sel_date_ai)
                if raw is not None and not pd.isna(raw):
                    try:
                        cell_value = float(raw)
                    except (ValueError, TypeError):
                        cell_value = None

            ref_info = reference_ranges.get(sel_test_ai, {})
            ref_min = ref_info.get('min') if isinstance(ref_info, dict) else None
            ref_max = ref_info.get('max') if isinstance(ref_info, dict) else None
            unit = units_map.get(sel_test_ai, '')
            age = patient_profile.get('age') if patient_profile else None
            gender = patient_gender or ''

            if cell_value is not None:
                ref_part = f" (ref: {ref_min}–{ref_max} {unit})" if ref_min is not None else ""
                st.info(f"**{sel_test_ai}** on {sel_date_ai}: **{cell_value} {unit}**{ref_part}")
            else:
                st.warning("No value recorded for this test on the selected date.")

            # ── Method selector ───────────────────────────────────────────────
            _hm_method_options = ["🌐 Web Search (no API key needed)"]
            if configured_svcs:
                _hm_method_options += [f"🤖 AI — {ai_mod.SERVICES[s]['label']}" for s in configured_svcs]

            hm_method_col, hm_btn_col = st.columns([4, 1])
            with hm_method_col:
                hm_chosen = st.selectbox(
                    "Lookup method",
                    options=_hm_method_options,
                    key=f"{key_prefix}_hm_method",
                    label_visibility="collapsed",
                )
            with hm_btn_col:
                hm_run = st.button(
                    "🔍 Analyse",
                    key=f"{key_prefix}_ai_run",
                    type="primary",
                    disabled=cell_value is None,
                )

            if hm_chosen.startswith("🌐"):
                st.caption("Searches DuckDuckGo and Wikipedia. No API key needed. Always verify results with a clinician.")
            else:
                st.caption("Uses your stored API key. Higher accuracy and medical context.")

            if hm_run:
                if cell_value is None:
                    st.warning("No value for the selected test/date combination.")
                else:
                    _hm_is_web = hm_chosen.startswith("🌐")
                    if _hm_is_web:
                        with st.spinner(f"Searching the web for information on {sel_test_ai}…"):
                            _hm_result, _hm_err, _hm_debug = ai_mod.scrape_test_explanation(
                                sel_test_ai, cell_value, unit, ref_min, ref_max, gender
                            )
                        if _hm_err:
                            st.error(f"Search failed: {_hm_err}")
                        else:
                            st.markdown(_hm_result)
                        if _hm_debug:
                            with st.expander("🔎 Debug — what was searched and what came back", expanded=_hm_err is not None):
                                for q in _hm_debug.get('queries', []):
                                    st.markdown(f"**{q.get('source', '?')}**")
                                    st.code(q.get('url', ''), language=None)
                                    if 'error' in q:
                                        st.error(f"Error: {q['error']}")
                                    else:
                                        st.markdown(f"HTTP {q.get('status','?')} · {q.get('chars',0)} chars")
                                        st.text(q.get('got', '(nothing)'))
                                    st.markdown("---")
                    else:
                        _hm_svc = None
                        for s in configured_svcs:
                            if ai_mod.SERVICES[s]['label'] in hm_chosen:
                                _hm_svc = s
                                break
                        if not _hm_svc:
                            st.warning("Could not identify the selected AI service.")
                        else:
                            prompt = ai_mod.build_test_prompt(
                                selected_patient, age, gender, sel_test_ai,
                                cell_value, unit, ref_min, ref_max,
                            )
                            with st.spinner(f"Querying {ai_mod.SERVICES[_hm_svc]['label']}…"):
                                response, err = ai_mod.query_service(_hm_svc, prompt, stored_keys[_hm_svc])
                            if err:
                                st.error(f"Error: {err}")
                            else:
                                st.markdown(response)

            # ── AI prompt viewer ──────────────────────────────────────────────
            prompt = ai_mod.build_test_prompt(
                selected_patient, age, gender, sel_test_ai,
                cell_value if cell_value is not None else "not recorded",
                unit, ref_min, ref_max,
            )
            with st.expander("📋 View AI prompt", expanded=False):
                st.text(prompt)

            # ── Open in Browser buttons ───────────────────────────────────────
            st.markdown("**Open in browser** *(use your existing AI login — prompt auto-copied where supported)*")
            open_cols = st.columns(len(ai_mod.SERVICES))
            for i, (svc, info) in enumerate(ai_mod.SERVICES.items()):
                url = ai_mod.get_quick_open_url(svc, prompt)
                note = " ✓ prompt auto-filled" if info["supports_url_query"] else ""
                with open_cols[i]:
                    st.markdown(
                        f'<a href="{url}" target="_blank" style="text-decoration:none;">'
                        f'<button style="background:{info["color"]};color:white;border:none;'
                        f'padding:7px 12px;border-radius:5px;font-size:13px;cursor:pointer;width:100%;">'
                        f'{info["icon"]} {svc}</button></a>',
                        unsafe_allow_html=True,
                    )
                    if note:
                        st.caption(note)

        def get_deviation_color_rgb(value, ref_info):
            if value is None or pd.isna(value):
                return None
            try:
                value = float(value)
            except (ValueError, TypeError):
                return None
            low, high = None, None
            if isinstance(ref_info, dict) and ref_info:
                low = ref_info.get('min')
                high = ref_info.get('max')
            if low is None or high is None:
                return None
            try:
                low, high = float(low), float(high)
            except (ValueError, TypeError):
                return None
            if high <= low:
                return None
            range_size = high - low
            if low <= value <= high:
                mid = (low + high) / 2
                deviation = abs(value - mid) / (range_size / 2)
                deviation = min(max(deviation, 0), 1)
                r = int(144 - deviation * 110)
                g = int(238 - deviation * 99)
                b = int(144 - deviation * 110)
                return (r, g, b)
            elif value < low:
                deviation = (low - value) / range_size
                deviation = min(max(deviation, 0), 2)
                intensity = min(deviation / 2, 1)
                r = int(173 - intensity * 173)
                g = int(216 - intensity * 216)
                b = int(230 - intensity * 91)
                return (r, g, b)
            else:
                deviation = (value - high) / range_size
                deviation = min(max(deviation, 0), 2)
                intensity = min(deviation / 2, 1)
                r = int(255 - intensity * 116)
                g = int(204 - intensity * 204)
                b = int(203 - intensity * 203)
                return (r, g, b)

        def get_tests_grouped_by_category():
            grouped = {}
            for test in unique_tests:
                info = test_info.get(test)
                cat = info['category'] if info else 'Other'
                if cat not in grouped:
                    grouped[cat] = []
                grouped[cat].append(test)
            ordered = []
            for cat in sorted(grouped.keys()):
                for t in sorted(grouped[cat], key=lambda x: x.lower()):
                    ordered.append(t)
            return ordered

        def generate_excel_heatmap():
            from openpyxl.worksheet.filters import AutoFilter
            wb = Workbook()
            ws = wb.active
            ws.title = "Heat Map"

            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF", size=9)
            summary_fill = PatternFill(start_color="E8E8E8", end_color="E8E8E8", fill_type="solid")
            summary_font = Font(bold=True, size=9)
            cat_separator_fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
            cat_separator_font = Font(bold=True, size=9, color="1F4E79")
            thin_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )
            center_align = Alignment(horizontal='center', vertical='center')
            rotated_align = Alignment(horizontal='center', vertical='bottom', text_rotation=90)

            excel_date_columns = list(unique_dates)
            excel_event_col_map = {}
            if show_events and event_dates_map:
                all_date_objs_excel = []
                for d in unique_dates:
                    try:
                        if date_format == "UK (DD/MM/YYYY)":
                            dt = pd.to_datetime(d, format='%d/%m/%Y')
                        else:
                            dt = pd.to_datetime(d, format='%m/%d/%Y')
                        all_date_objs_excel.append((dt, d, 'data'))
                    except:
                        all_date_objs_excel.append((pd.Timestamp.min, d, 'data'))
                for ev_formatted, ev_names in event_dates_map.items():
                    if ev_formatted not in unique_dates:
                        try:
                            if date_format == "UK (DD/MM/YYYY)":
                                ev_dt = pd.to_datetime(ev_formatted, format='%d/%m/%Y')
                            else:
                                ev_dt = pd.to_datetime(ev_formatted, format='%m/%d/%Y')
                            all_date_objs_excel.append((ev_dt, ev_formatted, 'event'))
                        except:
                            all_date_objs_excel.append((pd.Timestamp.min, ev_formatted, 'event'))
                all_date_objs_excel.sort(key=lambda x: x[0], reverse=True)
                excel_date_columns = []
                for dt_obj, formatted, dtype in all_date_objs_excel:
                    if dtype == 'event' and formatted not in unique_dates:
                        event_label = f"Event: {formatted}"
                        excel_date_columns.append(event_label)
                        excel_event_col_map[event_label] = event_dates_map.get(formatted, [])
                    else:
                        excel_date_columns.append(formatted)
                        if formatted in event_dates_map:
                            event_label = f"Event: {formatted}"
                            excel_date_columns.append(event_label)
                            excel_event_col_map[event_label] = event_dates_map.get(formatted, [])

            n_excel_date_cols = len(excel_date_columns)
            event_fill = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
            event_font = Font(bold=True, size=8, color="856404")

            headers = ['Test Name', 'Category', 'Range'] + excel_date_columns + ['Avg', 'Count', 'Sum']
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = thin_border
                if col_idx >= 4 and col_idx < 4 + n_excel_date_cols:
                    cell.alignment = rotated_align
                else:
                    cell.alignment = center_align

            ws.row_dimensions[1].height = 80

            grouped_tests = get_tests_grouped_by_category()
            current_row = 2
            current_category = None

            for test in grouped_tests:
                info = test_info.get(test)
                category = info['category'] if info else 'Other'

                if category != current_category:
                    current_category = category
                    total_cols = 3 + n_excel_date_cols + 3
                    cell = ws.cell(row=current_row, column=1, value=f"▸ {category}")
                    cell.font = cat_separator_font
                    cell.fill = cat_separator_fill
                    cell.border = thin_border
                    for c in range(2, total_cols + 1):
                        cell = ws.cell(row=current_row, column=c)
                        cell.fill = cat_separator_fill
                        cell.border = thin_border
                    current_row += 1

                range_str = range_column.get(test, 'N/A')

                ws.cell(row=current_row, column=1, value=test).border = thin_border
                ws.cell(row=current_row, column=1).font = Font(size=8)
                ws.cell(row=current_row, column=2, value=category).border = thin_border
                ws.cell(row=current_row, column=2).font = Font(size=8)
                ws.cell(row=current_row, column=3, value=range_str).border = thin_border
                ws.cell(row=current_row, column=3).alignment = center_align
                ws.cell(row=current_row, column=3).font = Font(size=8)

                ref_info = reference_ranges.get(test)
                for date_idx, date_col in enumerate(excel_date_columns):
                    col_num = 4 + date_idx
                    cell = ws.cell(row=current_row, column=col_num)
                    cell.border = thin_border
                    cell.alignment = center_align
                    if date_col in excel_event_col_map:
                        event_name_str = ' | '.join(excel_event_col_map[date_col])
                        cell.value = event_name_str
                        cell.fill = event_fill
                        cell.font = event_font
                    else:
                        val = pivot_data.get(test, {}).get(date_col)
                        if val is not None and not pd.isna(val):
                            try:
                                cell.value = round(float(val), 2)
                            except:
                                cell.value = val
                            cell.number_format = '0.0'
                            rgb = get_deviation_color_rgb(val, ref_info)
                            if rgb:
                                hex_color = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
                                cell.fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
                                lum = (0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2])
                                text_color = "FFFFFF" if lum < 128 else "000000"
                                cell.font = Font(color=text_color, size=8)

                stats = stats_data.get(test, {})
                stats_start = 4 + n_excel_date_cols
                for si, sv in enumerate([stats.get('Average', 0), stats.get('Count', 0), stats.get('Sum', 0)]):
                    cell = ws.cell(row=current_row, column=stats_start + si, value=sv)
                    cell.border = thin_border
                    cell.alignment = center_align
                    cell.font = Font(size=8)
                    if si != 1:
                        cell.number_format = '0.0'

                current_row += 1

            summary_row_num = current_row
            ws.cell(row=summary_row_num, column=1, value="Average Per Date").font = summary_font
            ws.cell(row=summary_row_num, column=1).fill = summary_fill
            ws.cell(row=summary_row_num, column=1).border = thin_border
            ws.cell(row=summary_row_num, column=2, value="").fill = summary_fill
            ws.cell(row=summary_row_num, column=2).border = thin_border
            ws.cell(row=summary_row_num, column=3, value="SUMMARY").fill = summary_fill
            ws.cell(row=summary_row_num, column=3).border = thin_border
            ws.cell(row=summary_row_num, column=3).alignment = center_align

            for date_idx, date_col in enumerate(excel_date_columns):
                col_num = 4 + date_idx
                cell = ws.cell(row=summary_row_num, column=col_num)
                cell.fill = summary_fill
                cell.font = summary_font
                cell.border = thin_border
                cell.alignment = center_align
                if date_col not in excel_event_col_map:
                    date_values = []
                    for test in unique_tests:
                        val = pivot_data.get(test, {}).get(date_col)
                        if val is not None and not pd.isna(val):
                            try:
                                date_values.append(float(val))
                            except:
                                pass
                    if date_values:
                        cell.value = round(sum(date_values) / len(date_values), 2)
                        cell.number_format = '0.0'

            stats_start = 4 + n_excel_date_cols
            for si, sv in enumerate([overall_avg, total_count, total_sum]):
                cell = ws.cell(row=summary_row_num, column=stats_start + si, value=round(sv, 2) if isinstance(sv, float) else sv)
                cell.fill = summary_fill
                cell.font = summary_font
                cell.border = thin_border
                cell.alignment = center_align

            ws.column_dimensions['A'].width = 18
            ws.column_dimensions['B'].width = 14
            ws.column_dimensions['C'].width = 12
            for col_idx in range(4, 4 + n_excel_date_cols):
                ws.column_dimensions[get_column_letter(col_idx)].width = 5.5
            for col_idx in range(4 + n_excel_date_cols, 4 + n_excel_date_cols + 3):
                ws.column_dimensions[get_column_letter(col_idx)].width = 6.5

            ws.freeze_panes = 'D2'

            last_col_letter = get_column_letter(3 + n_excel_date_cols + 3)
            ws.auto_filter.ref = f"A1:{last_col_letter}{summary_row_num}"

            legend_ws = wb.create_sheet("Color Legend")
            legend_ws.cell(row=1, column=1, value="Color Legend").font = Font(bold=True, size=14)
            legend_ws.cell(row=3, column=1, value="Below Normal Range").font = Font(bold=True, size=11)
            legend_ws.cell(row=3, column=2).fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")
            legend_ws.cell(row=3, column=3, value="Light Blue → Dark Blue (further below)")
            legend_ws.cell(row=4, column=2).fill = PatternFill(start_color="0000A0", end_color="0000A0", fill_type="solid")
            legend_ws.cell(row=5, column=1, value="Normal Range").font = Font(bold=True, size=11)
            legend_ws.cell(row=5, column=2).fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
            legend_ws.cell(row=5, column=3, value="Light Green → Dark Green (closer to boundary)")
            legend_ws.cell(row=6, column=2).fill = PatternFill(start_color="228B22", end_color="228B22", fill_type="solid")
            legend_ws.cell(row=7, column=1, value="Above Normal Range").font = Font(bold=True, size=11)
            legend_ws.cell(row=7, column=2).fill = PatternFill(start_color="FFCCCB", end_color="FFCCCB", fill_type="solid")
            legend_ws.cell(row=7, column=3, value="Light Red → Dark Red (further above)")
            legend_ws.cell(row=8, column=2).fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
            legend_ws.column_dimensions['A'].width = 25
            legend_ws.column_dimensions['B'].width = 10
            legend_ws.column_dimensions['C'].width = 45
            if patient_gender:
                legend_ws.cell(row=10, column=1, value=f"Reference ranges: {patient_gender.title()}").font = Font(italic=True, size=11)

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output

        def generate_heatmap_image():
            grouped_tests = get_tests_grouped_by_category()
            categories_in_order = []
            seen_cats = set()
            for test in grouped_tests:
                info = test_info.get(test)
                cat = info['category'] if info else 'Other'
                if cat not in seen_cats:
                    categories_in_order.append(cat)
                    seen_cats.add(cat)

            img_date_columns = list(unique_dates)
            img_event_col_map = {}
            if show_events and event_dates_map:
                all_date_objs_img = []
                for d in unique_dates:
                    try:
                        if date_format == "UK (DD/MM/YYYY)":
                            dt = pd.to_datetime(d, format='%d/%m/%Y')
                        else:
                            dt = pd.to_datetime(d, format='%m/%d/%Y')
                        all_date_objs_img.append((dt, d, 'data'))
                    except:
                        all_date_objs_img.append((pd.Timestamp.min, d, 'data'))
                for ev_formatted, ev_names in event_dates_map.items():
                    if ev_formatted not in unique_dates:
                        try:
                            if date_format == "UK (DD/MM/YYYY)":
                                ev_dt = pd.to_datetime(ev_formatted, format='%d/%m/%Y')
                            else:
                                ev_dt = pd.to_datetime(ev_formatted, format='%m/%d/%Y')
                            all_date_objs_img.append((ev_dt, ev_formatted, 'event'))
                        except:
                            all_date_objs_img.append((pd.Timestamp.min, ev_formatted, 'event'))
                all_date_objs_img.sort(key=lambda x: x[0], reverse=True)
                img_date_columns = []
                for dt_obj, formatted, dtype in all_date_objs_img:
                    if dtype == 'event' and formatted not in unique_dates:
                        event_label = f"Ev: {formatted}"
                        img_date_columns.append(event_label)
                        img_event_col_map[event_label] = event_dates_map.get(formatted, [])
                    else:
                        img_date_columns.append(formatted)
                        if formatted in event_dates_map:
                            event_label = f"Ev: {formatted}"
                            img_date_columns.append(event_label)
                            img_event_col_map[event_label] = event_dates_map.get(formatted, [])

            n_cat_separators = len(categories_in_order)
            n_data_rows = len(unique_tests)
            n_rows = n_data_rows + n_cat_separators + 1
            n_img_date_cols = len(img_date_columns)
            n_stat_cols = 3
            fixed_cols = ['Test Name', 'Range']
            n_fixed = len(fixed_cols)
            total_cols = n_fixed + n_img_date_cols + n_stat_cols

            col_w_fixed = [2.0, 1.2]
            col_w_date = [0.50] * n_img_date_cols
            col_w_stats = [0.55, 0.45, 0.55]
            col_widths = col_w_fixed + col_w_date + col_w_stats
            fig_width = sum(col_widths) + 0.3
            row_height = 0.26
            header_height = 1.8
            fig_height = header_height + n_rows * row_height + 0.5

            fig, ax = plt.subplots(figsize=(fig_width, fig_height), dpi=200)
            ax.set_xlim(0, sum(col_widths))
            ax.set_ylim(0, header_height + n_rows * row_height)
            ax.axis('off')
            fig.patch.set_facecolor('white')

            x_positions = [0]
            for w in col_widths:
                x_positions.append(x_positions[-1] + w)

            total_height = header_height + n_rows * row_height

            header_color = '#4472C4'
            header_y = total_height - header_height
            ax.add_patch(plt.Rectangle((0, header_y), sum(col_widths), header_height, facecolor=header_color, edgecolor='#D9D9D9', linewidth=0.5))

            all_headers = fixed_cols + img_date_columns + ['Avg', 'Count', 'Sum']
            for col_idx, header in enumerate(all_headers):
                cx = x_positions[col_idx] + col_widths[col_idx] / 2
                if col_idx >= n_fixed and col_idx < n_fixed + n_img_date_cols:
                    h_color = '#856404' if header in img_event_col_map else 'white'
                    ax.text(cx, header_y + 0.15, header, ha='center', va='bottom',
                            fontsize=4.5, fontweight='bold', color=h_color, rotation=90,
                            fontfamily='sans-serif')
                else:
                    cy = header_y + header_height / 2
                    ax.text(cx, cy, header, ha='center', va='center',
                            fontsize=5 if col_idx < n_fixed else 4.5, fontweight='bold', color='white',
                            fontfamily='sans-serif')

            visual_row = 0
            current_category = None
            cat_sep_color = '#D6E4F0'
            cat_sep_text_color = '#1F4E79'

            for test in grouped_tests:
                info_t = test_info.get(test)
                category = info_t['category'] if info_t else 'Other'

                if category != current_category:
                    current_category = category
                    y = total_height - header_height - (visual_row + 1) * row_height
                    ax.add_patch(plt.Rectangle((0, y), sum(col_widths), row_height, facecolor=cat_sep_color, edgecolor='#D9D9D9', linewidth=0.3))
                    ax.text(x_positions[0] + 0.06, y + row_height / 2, f"▸ {category}", ha='left', va='center',
                            fontsize=4.5, fontweight='bold', color=cat_sep_text_color, fontfamily='sans-serif')
                    visual_row += 1

                y = total_height - header_height - (visual_row + 1) * row_height
                row_bg = '#FFFFFF' if visual_row % 2 == 0 else '#F8F8F8'
                ax.add_patch(plt.Rectangle((0, y), sum(col_widths), row_height, facecolor=row_bg, edgecolor='#D9D9D9', linewidth=0.3))

                range_str = range_column.get(test, 'N/A')
                ax.text(x_positions[0] + 0.06, y + row_height / 2, test, ha='left', va='center', fontsize=4, fontfamily='sans-serif', clip_on=True)
                ax.text(x_positions[1] + col_widths[1] / 2, y + row_height / 2, range_str, ha='center', va='center', fontsize=3.5, fontfamily='sans-serif')

                ref_info = reference_ranges.get(test)
                for date_idx, date_col in enumerate(img_date_columns):
                    col_num = n_fixed + date_idx
                    cx = x_positions[col_num] + col_widths[col_num] / 2
                    cell_x = x_positions[col_num]
                    cell_w = col_widths[col_num]

                    if date_col in img_event_col_map:
                        event_name_str = ' | '.join(img_event_col_map[date_col])
                        ax.add_patch(plt.Rectangle((cell_x + 0.02, y + 0.02), cell_w - 0.04, row_height - 0.04,
                            facecolor='#FFF3CD', edgecolor='none', linewidth=0))
                        ax.text(cx, y + row_height / 2, event_name_str[:8], ha='center', va='center', fontsize=2.5, color='#856404', fontfamily='sans-serif', clip_on=True)
                    else:
                        val = pivot_data.get(test, {}).get(date_col)
                        if val is not None and not pd.isna(val):
                            try:
                                fval = float(val)
                            except:
                                continue
                            rgb = get_deviation_color_rgb(fval, ref_info)
                            if rgb:
                                norm_color = (rgb[0]/255, rgb[1]/255, rgb[2]/255)
                                ax.add_patch(plt.Rectangle((cell_x + 0.02, y + 0.02), cell_w - 0.04, row_height - 0.04,
                                    facecolor=norm_color, edgecolor='none', linewidth=0))
                                lum = 0.299 * rgb[0] + 0.587 * rgb[1] + 0.114 * rgb[2]
                                tc = 'white' if lum < 128 else 'black'
                            else:
                                tc = 'black'
                            if fval >= 1000:
                                display = f"{fval:.0f}"
                            elif fval >= 100:
                                display = f"{fval:.1f}"
                            else:
                                display = f"{fval:.2f}"
                            ax.text(cx, y + row_height / 2, display, ha='center', va='center', fontsize=3.2, color=tc, fontfamily='sans-serif')

                stats = stats_data.get(test, {})
                stat_values = [stats.get('Average', 0), stats.get('Count', 0), stats.get('Sum', 0)]
                for si, sv in enumerate(stat_values):
                    col_num = n_fixed + n_img_date_cols + si
                    cx = x_positions[col_num] + col_widths[col_num] / 2
                    if isinstance(sv, float):
                        display = f"{sv:.1f}" if sv >= 100 else f"{sv:.2f}"
                    else:
                        display = str(sv)
                    ax.text(cx, y + row_height / 2, display, ha='center', va='center', fontsize=3.5, fontfamily='sans-serif')

                visual_row += 1

            y = total_height - header_height - (visual_row + 1) * row_height
            ax.add_patch(plt.Rectangle((0, y), sum(col_widths), row_height, facecolor='#E8E8E8', edgecolor='#D9D9D9', linewidth=0.3))
            ax.text(x_positions[0] + 0.06, y + row_height / 2, 'Avg Per Date', ha='left', va='center', fontsize=4, fontweight='bold', fontfamily='sans-serif')
            ax.text(x_positions[1] + col_widths[1] / 2, y + row_height / 2, 'SUMMARY', ha='center', va='center', fontsize=3.5, fontweight='bold', fontfamily='sans-serif')

            for date_idx, date_col in enumerate(img_date_columns):
                col_num = n_fixed + date_idx
                cx = x_positions[col_num] + col_widths[col_num] / 2
                if date_col in img_event_col_map:
                    continue
                date_values = []
                for t in unique_tests:
                    val = pivot_data.get(t, {}).get(date_col)
                    if val is not None and not pd.isna(val):
                        try:
                            date_values.append(float(val))
                        except:
                            pass
                if date_values:
                    avg_val = sum(date_values) / len(date_values)
                    display = f"{avg_val:.1f}" if avg_val >= 100 else f"{avg_val:.2f}"
                    ax.text(cx, y + row_height / 2, display, ha='center', va='center', fontsize=3.2, fontweight='bold', fontfamily='sans-serif')

            summary_stats = [overall_avg, total_count, total_sum]
            for si, sv in enumerate(summary_stats):
                col_num = n_fixed + n_img_date_cols + si
                cx = x_positions[col_num] + col_widths[col_num] / 2
                if isinstance(sv, float):
                    display = f"{sv:.1f}" if sv >= 100 else f"{sv:.2f}"
                else:
                    display = str(sv)
                ax.text(cx, y + row_height / 2, display, ha='center', va='center', fontsize=3.5, fontweight='bold', fontfamily='sans-serif')

            for col_idx in range(total_cols + 1):
                x = x_positions[col_idx] if col_idx < len(x_positions) else sum(col_widths)
                ax.plot([x, x], [0, total_height], color='#D9D9D9', linewidth=0.3)
            for r in range(n_rows + 1):
                y_line = total_height - header_height - r * row_height
                ax.plot([0, sum(col_widths)], [y_line, y_line], color='#D9D9D9', linewidth=0.3)

            legend_y = -0.15
            legend_items = [
                ('Below Normal', '#ADD8E6', '#00008B'),
                ('Normal Range', '#90EE90', '#228B22'),
                ('Above Normal', '#FFCCCB', '#8B0000')
            ]
            legend_x = 0.3
            for label, c1, c2 in legend_items:
                ax.text(legend_x, legend_y, label + ': ', ha='left', va='center', fontsize=4.5, fontweight='bold', fontfamily='sans-serif', transform=ax.transData)
                legend_x += 1.3
                for i in range(20):
                    t = i / 19
                    rc1 = mcolors.to_rgb(c1)
                    rc2 = mcolors.to_rgb(c2)
                    mixed = tuple(rc1[j] * (1-t) + rc2[j] * t for j in range(3))
                    ax.add_patch(plt.Rectangle((legend_x + i * 0.05, legend_y - 0.06), 0.05, 0.12, facecolor=mixed, edgecolor='none'))
                legend_x += 1.3 + 0.3

            if patient_gender:
                ax.text(sum(col_widths) - 0.1, legend_y, f"Ref ranges: {patient_gender.title()}", ha='right', va='center', fontsize=4, fontstyle='italic', fontfamily='sans-serif')

            ax.text(sum(col_widths) / 2, total_height + 0.15, f"Heat Map - {selected_patient} ({len(unique_tests)} tests × {len(unique_dates)} dates)",
                    ha='center', va='bottom', fontsize=7, fontweight='bold', fontfamily='sans-serif')

            plt.subplots_adjust(left=0.01, right=0.99, top=0.97, bottom=0.04)

            buf = io.BytesIO()
            fig.savefig(buf, format='png', dpi=300, bbox_inches='tight', facecolor='white', edgecolor='none')
            plt.close(fig)
            buf.seek(0)
            return buf

        dl_col1, dl_col2, dl_col3 = st.columns([1, 1, 2])
        with dl_col1:
            excel_data = generate_excel_heatmap()
            safe_name = selected_patient.replace(' ', '_')
            st.download_button(
                label="📥 Download as Excel",
                data=excel_data,
                file_name=f"heatmap_{safe_name}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"{key_prefix}_download_heatmap_excel"
            )
        with dl_col2:
            image_data = generate_heatmap_image()
            st.download_button(
                label="🖼️ Download as Image",
                data=image_data,
                file_name=f"heatmap_{safe_name}.png",
                mime="image/png",
                key=f"{key_prefix}_download_heatmap_image"
            )

        if editable:
            with st.expander("✏️ Editable Table", expanded=True):
                st.caption("Edit values directly. Clear a cell to delete that value. Stats columns (Avg, Count, Sum) are read-only. Test Name column stays frozen when scrolling.")

                added_dates_key = f"{key_prefix}_added_dates_{selected_patient}"
                added_tests_key = f"{key_prefix}_added_tests_{selected_patient}"
                if added_dates_key not in st.session_state:
                    st.session_state[added_dates_key] = []
                if added_tests_key not in st.session_state:
                    st.session_state[added_tests_key] = []

                for ad in st.session_state[added_dates_key]:
                    if ad not in unique_dates:
                        unique_dates.append(ad)
                        for test in unique_tests:
                            if ad not in pivot_data.get(test, {}):
                                pivot_data[test][ad] = None
                                value_ids[test][ad] = None
                for at in st.session_state[added_tests_key]:
                    if at not in unique_tests:
                        unique_tests.append(at)
                        pivot_data[at] = {d: None for d in unique_dates}
                        value_ids[at] = {d: None for d in unique_dates}

                add_col1, add_col2 = st.columns(2)
                with add_col1:
                    with st.popover("➕ Add New Date Column"):
                        new_date = st.date_input("New date:", key=f"{key_prefix}_new_date_input")
                        if st.button("Add Date", key=f"{key_prefix}_add_date_btn"):
                            if date_format == "UK (DD/MM/YYYY)":
                                new_date_str = new_date.strftime('%d/%m/%Y')
                            else:
                                new_date_str = new_date.strftime('%m/%d/%Y')
                            if new_date_str not in unique_dates:
                                st.session_state[added_dates_key].append(new_date_str)
                                st.success(f"Added date column: {new_date_str}")
                                editor_key = f"{key_prefix}_heatmap_editor"
                                if editor_key in st.session_state:
                                    del st.session_state[editor_key]
                                st.rerun()
                            else:
                                st.warning("This date already exists.")
                with add_col2:
                    with st.popover("➕ Add New Test Row"):
                        all_standard_names = sorted([s.get('standard_name', '') for s in all_standards if isinstance(s, dict) and s.get('standard_name')], key=str.lower)
                        existing_tests = set(unique_tests)
                        available_tests = [t for t in all_standard_names if t not in existing_tests]
                        if available_tests:
                            new_test = st.selectbox("Select test from standards:", available_tests, index=None, placeholder="Choose a test...", key=f"{key_prefix}_new_test_select")
                        else:
                            new_test = None
                            st.info("All standard tests are already in the table.")
                        custom_test = st.text_input("Or enter custom test name:", key=f"{key_prefix}_custom_test_input")
                        if st.button("Add Test", key=f"{key_prefix}_add_test_btn"):
                            test_to_add = custom_test.strip() if custom_test.strip() else new_test
                            if test_to_add and test_to_add not in existing_tests:
                                st.session_state[added_tests_key].append(test_to_add)
                                st.success(f"Added test row: {test_to_add}")
                                editor_key = f"{key_prefix}_heatmap_editor"
                                if editor_key in st.session_state:
                                    del st.session_state[editor_key]
                                st.rerun()
                            elif test_to_add in existing_tests:
                                st.warning("This test already exists in the table.")
                            else:
                                st.warning("Please select or enter a test name.")

                editable_df = pd.DataFrame(pivot_data).T
                editable_df.index.name = 'Test Name'
                for d in unique_dates:
                    if d not in editable_df.columns:
                        editable_df[d] = None
                editable_df = editable_df[unique_dates]

                editable_df.insert(0, 'Range', [range_column.get(t, 'N/A') for t in editable_df.index])
                editable_df['Avg'] = [stats_data.get(t, {}).get('Average', 0) for t in editable_df.index]
                editable_df['Count'] = [stats_data.get(t, {}).get('Count', 0) for t in editable_df.index]
                editable_df['Sum'] = [stats_data.get(t, {}).get('Sum', 0) for t in editable_df.index]
                editable_df = editable_df.reset_index()
                editable_df = editable_df.rename(columns={'index': 'Test Name'})

                for date_col in unique_dates:
                    if date_col in editable_df.columns:
                        editable_df[date_col] = editable_df[date_col].astype(object)
                        editable_df[date_col] = editable_df[date_col].where(editable_df[date_col].notna(), None)

                column_config = {
                    'Test Name': st.column_config.TextColumn('Test Name', disabled=True, pinned=True),
                    'Range': st.column_config.TextColumn('Range', disabled=True),
                    'Avg': st.column_config.NumberColumn('Avg', disabled=True, format="%.2f"),
                    'Count': st.column_config.NumberColumn('Count', disabled=True),
                    'Sum': st.column_config.NumberColumn('Sum', disabled=True, format="%.2f")
                }
                for date_col in unique_dates:
                    column_config[date_col] = st.column_config.NumberColumn(date_col, format="%.2f", required=False)

                edited_df = st.data_editor(
                    editable_df,
                    use_container_width=False,
                    hide_index=True,
                    num_rows="fixed",
                    column_config=column_config,
                    key=f"{key_prefix}_heatmap_editor"
                )

            col_save1, col_save2 = st.columns([1, 3])
            with col_save1:
                save_btn = st.button("💾 Save All Changes", key=f"{key_prefix}_save_heatmap_changes", type="primary")
            with col_save2:
                if st.button("🔄 Refresh Data", key=f"{key_prefix}_refresh_heatmap"):
                    editor_key = f"{key_prefix}_heatmap_editor"
                    if editor_key in st.session_state:
                        del st.session_state[editor_key]
                    st.rerun()

            if save_btn:
                changes_made = 0
                errors = 0
                new_entries = 0
                deleted_entries = 0
                changes_list = []

                def _is_empty(v):
                    if v is None:
                        return True
                    if isinstance(v, str):
                        return v.strip() == ''
                    try:
                        return pd.isna(v)
                    except (TypeError, ValueError):
                        return False

                editor_state = st.session_state.get(f"{key_prefix}_heatmap_editor", {})
                edited_rows = editor_state.get('edited_rows', {}) if isinstance(editor_state, dict) else {}

                test_name_list = list(editable_df['Test Name'])

                cells_to_save = []

                for row_idx_str, col_changes in edited_rows.items():
                    row_idx = int(row_idx_str)
                    if row_idx < 0 or row_idx >= len(test_name_list):
                        continue
                    test_name = test_name_list[row_idx]
                    for col_name, new_value in col_changes.items():
                        if col_name in unique_dates:
                            original_value = pivot_data.get(test_name, {}).get(col_name)
                            result_id = value_ids.get(test_name, {}).get(col_name)
                            cells_to_save.append((test_name, col_name, new_value, original_value, result_id))

                for idx, row in edited_df.iterrows():
                    test_name = row['Test Name']
                    for date_col in unique_dates:
                        if date_col not in row.index:
                            continue
                        new_value = row[date_col]
                        original_value = pivot_data.get(test_name, {}).get(date_col)
                        result_id = value_ids.get(test_name, {}).get(date_col)
                        new_is_empty = _is_empty(new_value)
                        orig_is_empty = _is_empty(original_value)
                        if new_is_empty and orig_is_empty:
                            continue
                        value_changed = False
                        if new_is_empty != orig_is_empty:
                            value_changed = True
                        elif not new_is_empty and not orig_is_empty:
                            try:
                                if abs(float(new_value) - float(original_value)) > 0.001:
                                    value_changed = True
                            except (ValueError, TypeError):
                                if str(new_value) != str(original_value):
                                    value_changed = True
                        if value_changed:
                            already_tracked = any(c[0] == test_name and c[1] == date_col for c in cells_to_save)
                            if not already_tracked:
                                cells_to_save.append((test_name, date_col, new_value, original_value, result_id))

                with st.expander("🔍 Changes Detected", expanded=False):
                    if cells_to_save:
                        for test_name, date_col, new_value, original_value, result_id in cells_to_save:
                            new_is_empty = _is_empty(new_value)
                            orig_is_empty = _is_empty(original_value)
                            if new_is_empty and not orig_is_empty:
                                st.info(f"🗑️ DELETE: **{test_name}** @ {date_col}: {original_value} → (empty)")
                            elif not new_is_empty and orig_is_empty:
                                st.success(f"➕ NEW: **{test_name}** @ {date_col}: → {new_value}")
                            elif not new_is_empty and not orig_is_empty:
                                st.success(f"✏️ UPDATE: **{test_name}** @ {date_col}: {original_value} → {new_value}")
                    else:
                        st.info("No changes detected")

                for test_name, date_col, new_value, original_value, result_id in cells_to_save:
                    new_is_empty = _is_empty(new_value)
                    orig_is_empty = _is_empty(original_value)

                    if new_is_empty and orig_is_empty:
                        continue

                    try:
                        try:
                            if date_format == "UK (DD/MM/YYYY)":
                                parsed_date = pd.to_datetime(date_col, format='%d/%m/%Y').date()
                            else:
                                parsed_date = pd.to_datetime(date_col, format='%m/%d/%Y').date()
                        except:
                            try:
                                parsed_date = pd.to_datetime(date_col).date()
                            except Exception as date_err:
                                st.error(f"Date parse error for {date_col}: {date_err}")
                                errors += 1
                                continue

                        unit = units_map.get(test_name, '')
                        ref_info = reference_ranges.get(test_name, '')
                        ref_str = f"{ref_info['min']}-{ref_info['max']}" if isinstance(ref_info, dict) else str(ref_info)

                        numeric_value = None
                        if not new_is_empty:
                            try:
                                numeric_value = float(new_value)
                            except (ValueError, TypeError):
                                errors += 1
                                st.error(f"Invalid number '{new_value}' for {test_name}")
                                continue

                        if result_id is not None:
                            if new_is_empty:
                                if db.delete_test_result(int(result_id)):
                                    deleted_entries += 1
                                    changes_list.append(f"{test_name} ({date_col}): {original_value} → deleted")
                                else:
                                    errors += 1
                                    st.error(f"Failed to delete {test_name} @ {date_col}")
                            else:
                                update_result = db.update_test_result(
                                    result_id=result_id,
                                    test_type=test_name,
                                    value=numeric_value,
                                    unit=unit,
                                    test_date=parsed_date,
                                    reference_range=ref_str
                                )
                                if update_result:
                                    changes_made += 1
                                    changes_list.append(f"{test_name} ({date_col}): {original_value} → {numeric_value}")
                                else:
                                    st.error(f"Update failed for {test_name} (id={result_id})")
                        else:
                            if not new_is_empty:
                                patient_id = db.get_patient_id(selected_patient, user_id)
                                if not patient_id:
                                    st.error(f"Could not find patient '{selected_patient}' - cannot add new test result")
                                    errors += 1
                                elif db.add_test_result(
                                    patient_id=patient_id,
                                    test_date=parsed_date,
                                    test_type=test_name,
                                    test_value=numeric_value,
                                    unit=unit,
                                    reference_range=ref_str,
                                    user_id=user_id
                                ):
                                    new_entries += 1
                                    changes_list.append(f"{test_name} ({date_col}): → {numeric_value}")
                                else:
                                    st.error(f"Failed to add {test_name}")
                                    errors += 1
                    except Exception as e:
                        errors += 1
                        st.error(f"Error saving {test_name}: {e}")

                total_changes = changes_made + new_entries + deleted_entries
                if total_changes > 0:
                    msg_parts = []
                    if changes_made > 0:
                        msg_parts.append(f"{changes_made} updated")
                    if new_entries > 0:
                        msg_parts.append(f"{new_entries} added")
                    if deleted_entries > 0:
                        msg_parts.append(f"{deleted_entries} deleted")
                    st.success(f"✅ {total_changes} change(s) saved: {', '.join(msg_parts)}")
                    if changes_list and len(changes_list) <= 10:
                        with st.expander("View changes"):
                            for change in changes_list:
                                st.write(f"• {change}")
                    elif changes_list:
                        with st.expander(f"View first 10 of {len(changes_list)} changes"):
                            for change in changes_list[:10]:
                                st.write(f"• {change}")
                    editor_key = f"{key_prefix}_heatmap_editor"
                    if editor_key in st.session_state:
                        del st.session_state[editor_key]
                    added_dates_key = f"{key_prefix}_added_dates_{selected_patient}"
                    added_tests_key = f"{key_prefix}_added_tests_{selected_patient}"
                    if added_dates_key in st.session_state:
                        del st.session_state[added_dates_key]
                    if added_tests_key in st.session_state:
                        del st.session_state[added_tests_key]
                    st.rerun()
                elif errors > 0:
                    st.error(f"❌ {errors} errors occurred while saving")
                else:
                    st.info("No changes detected. Edit values in the table above and click Save.")
